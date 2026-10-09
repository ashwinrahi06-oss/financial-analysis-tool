
"""Generate forecast assumptions and accounting reconciliation workbook."""

import argparse

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .model_config import (
    DATA,
    FILING_URL,
    FY26_DEPRECIATION,
    FY26_DA_AND_OTHER,
    FY26_FINANCE_LEASE_ADDITIONS,
    FY26_FINANCE_LEASE_LIABILITIES,
    FY26_OPERATING_LEASE_LIABILITIES,
    GROWTH_DECAY,
    MARGIN_DELTAS,
    CAPEX_RATIO_MULTIPLIERS,
    DA_RATIO_MULTIPLIERS,
    TAX_RATE,
    NWC_RATIO,
    FORECAST_YEARS,
)

from .financial_data import fetch_facts, load_financials


OUTPUT = DATA / "Microsoft_Forecast_Assumptions.xlsx"


def forecast_table(history, da_input=FY26_DEPRECIATION):
    revenue_2025 = float(
        history.loc[2025, "Revenue"]
    )

    revenue_2026 = float(
        history.loc[2026, "Revenue"]
    )

    historical_growth = (
        revenue_2026 / revenue_2025 - 1
    )

    historical_margin = (
        float(history.loc[2026, "Operating Income"])
        / revenue_2026
    )

    historical_capex = (
        float(history.loc[2026, "Cash CapEx"])
        / revenue_2026
    )

    historical_da = (
        float(da_input) / revenue_2026
    )

    rows = []

    for i, year in enumerate(FORECAST_YEARS):
        rows.append({
            "Fiscal Year": year,
            "Revenue Growth": (
                historical_growth * GROWTH_DECAY[i]
            ),
            "EBIT Margin": (
                historical_margin + MARGIN_DELTAS[i]
            ),
            "CapEx / Revenue": (
                historical_capex
                * CAPEX_RATIO_MULTIPLIERS[i]
            ),
            "D&A / Revenue": (
                historical_da
                * DA_RATIO_MULTIPLIERS[i]
            ),
            "Tax Rate": TAX_RATE,
            "Incremental NWC Ratio": NWC_RATIO,
        })

    return pd.DataFrame(rows)


def format_sheet(
    ws,
    widths=None,
    percent_cols=(),
):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34

    for cell in ws[1]:
        cell.fill = PatternFill(
            "solid",
            fgColor="17365D",
        )

        cell.font = Font(
            name="Aptos",
            bold=True,
            color="FFFFFF",
            size=11,
        )

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(
                name="Aptos",
                size=10,
                color="243247",
            )

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if row[0].row % 2 == 0:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="EDF3FA",
                )

            if (
                isinstance(cell.value, (int, float))
                and not isinstance(cell.value, bool)
            ):
                if cell.column in percent_cols:
                    cell.number_format = "0.00%"
                else:
                    cell.number_format = "#,##0.0"

    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[
            get_column_letter(col)
        ].width = 23

    for column, width in (widths or {}).items():
        ws.column_dimensions[column].width = width


def write_forecast(
    history,
    da_input=FY26_DEPRECIATION,
    da_status="DEPRECIATION-ONLY FLOOR",
):
    forecast = forecast_table(
        history,
        da_input,
    )

    wb = Workbook()

    # Forecast Drivers
    ws = wb.active
    ws.title = "Forecast Drivers"

    ws.append(list(forecast.columns))

    for row in forecast.itertuples(
        index=False,
        name=None,
    ):
        ws.append(list(row))

    format_sheet(
        ws,
        {"A": 18},
        percent_cols=range(2, 8),
    )

    ws.sheet_properties.tabColor = "2878B5"

    # Historical Drivers
    hist = wb.create_sheet("Historical Drivers")

    hist.append([
        "Metric",
        "FY2025",
        "FY2026",
        "Source",
    ])

    metrics = [
        "Revenue",
        "Operating Income",
        "Net Income",
        "Operating Cash Flow",
        "Cash CapEx",
        "Free Cash Flow",
    ]

    for metric in metrics:
        hist.append([
            metric,
            float(history.loc[2025, metric]),
            float(history.loc[2026, metric]),
            FILING_URL,
        ])

    format_sheet(
        hist,
        {"A": 30, "D": 80},
    )

    # CapEx and D&A Reconciliation
    recon = wb.create_sheet(
        "CapEx DA Reconciliation"
    )

    recon.append([
        "Item",
        "USD millions",
        "Classification",
        "Source / Caveat",
    ])

    records = [
        (
            "FY2026 cash PP&E purchases",
            float(history.loc[2026, "Cash CapEx"]),
            "Cash investing outflow",
            "SEC Company Facts / 10-K",
        ),
        (
            "FY2026 depreciation expense",
            FY26_DEPRECIATION,
            "Reported depreciation",
            "Not necessarily complete D&A",
        ),
        (
            "FY2026 D&A and other",
            FY26_DA_AND_OTHER,
            "Broader noncash adjustment",
            "Includes other items; not pure D&A",
        ),
        (
            "FY2026 D&A model input",
            float(da_input),
            da_status,
            "Explicit analyst assumption",
        ),
        (
            "New FY2026 finance lease ROU assets",
            FY26_FINANCE_LEASE_ADDITIONS,
            "Noncash lease additions",
            "10-K lease disclosures",
        ),
        (
            "FY2026 finance lease liabilities",
            FY26_FINANCE_LEASE_LIABILITIES,
            "Lease obligation",
            "Disclosed separately",
        ),
        (
            "FY2026 operating lease liabilities",
            FY26_OPERATING_LEASE_LIABILITIES,
            "Lease obligation",
            "Disclosed separately",
        ),
    ]

    for record in records:
        recon.append(record)

    format_sheet(
        recon,
        {
            "A": 44,
            "B": 20,
            "C": 36,
            "D": 85,
        },
    )

    for row in recon.iter_rows(min_row=2):
        if "D&A model input" in str(row[0].value):
            for cell in row:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="FFF2CC",
                )

    # Methodology
    methodology = wb.create_sheet("Methodology")

    methodology.append([
        "Forecast Methodology and Limitations"
    ])

    notes = [
        "Revenue growth moderates from FY2026 actual growth.",
        "EBIT margin forecasts are illustrative analyst assumptions.",
        "Cash CapEx forecasts use FY2026 cash PP&E purchases.",
        "Finance lease additions are not included in cash CapEx.",
        "The D&A starting value is a depreciation-only floor.",
        "Depreciation and amortization require further reconciliation.",
        "D&A intensity rises initially as infrastructure depreciates.",
        "Lease liabilities are disclosed separately from financial debt.",
        "WACC and terminal growth assumptions are editable.",
        "Forecasts are not Microsoft management guidance.",
        "Source: " + FILING_URL,
    ]

    for note in notes:
        methodology.append([note])

    format_sheet(
        methodology,
        {"A": 115},
    )

    wb.save(OUTPUT)

    return forecast, OUTPUT


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--da",
        type=float,
        default=FY26_DEPRECIATION,
        help="FY2026 D&A assumption in USD millions",
    )

    args = parser.parse_args()

    if args.da <= 0:
        raise ValueError(
            "D&A assumption must be positive."
        )

    history, _, _ = load_financials(
        fetch_facts()
    )

    status = (
        "DEPRECIATION-ONLY FLOOR"
        if args.da == FY26_DEPRECIATION
        else "USER-SUPPLIED / VERIFY"
    )

    forecast, path = write_forecast(
        history,
        args.da,
        status,
    )

    print(
        forecast.to_string(
            index=False,
            float_format=lambda value: f"{value:.4f}",
        )
    )

    print("Saved:", path)
    print("D&A status:", status)


if __name__ == "__main__":
    main()
