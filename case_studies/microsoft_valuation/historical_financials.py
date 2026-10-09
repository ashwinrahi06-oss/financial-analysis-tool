
from pathlib import Path
from datetime import date
import time

import pandas as pd
import requests
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT = DATA_DIR / "Microsoft_Historical_Validated.xlsx"

CIK = "0000789019"
COMPANY = "Microsoft Corporation"
YEARS = range(2022, 2027)

# IMPORTANT: Use your actual email address.
# The SEC asks automated clients to identify themselves.
USER_AGENT = "FinancialResearchStudent ashwinrahi06@gmail.com"

SEC_URL = (
    f"https://data.sec.gov/api/xbrl/companyfacts/CIK{CIK}.json"
)

TAGS = {
    "Revenue": [
        "RevenueFromContractWithCustomerExcludingAssessedTax",
        "Revenues",
        "SalesRevenueNet",
    ],
    "Operating Income": [
        "OperatingIncomeLoss",
    ],
    "Net Income": [
        "NetIncomeLoss",
    ],
    "Operating Cash Flow": [
        "NetCashProvidedByUsedInOperatingActivities",
    ],
    "Capital Expenditures": [
        "PaymentsToAcquirePropertyPlantAndEquipment",
    ],
    "Depreciation and Amortization": [
        "DepreciationDepletionAndAmortization",
        "DepreciationDepletionAndAmortizationPropertyPlantAndEquipment",
    ],
}


# ============================================================
# SEC DATA DOWNLOAD
# ============================================================

def download_company_facts():
    if "example.com" in USER_AGENT:
        raise ValueError(
            "Replace your-email@example.com with your actual "
            "email address in USER_AGENT before running."
        )

    headers = {
        "User-Agent": USER_AGENT,
        "Accept-Encoding": "gzip, deflate",
        "Host": "data.sec.gov",
    }

    response = requests.get(
        SEC_URL,
        headers=headers,
        timeout=30,
    )
    response.raise_for_status()
    return response.json()


# ============================================================
# ANNUAL FINANCIAL DATA EXTRACTION
# ============================================================

def get_annual_fact(facts, tags, year):
    us_gaap = facts.get("facts", {}).get("us-gaap", {})

    # Microsoft's fiscal year ends June 30.
    fiscal_end = f"{year}-06-30"

    candidates = []

    for priority, tag in enumerate(tags):
        concept = us_gaap.get(tag)

        if not concept:
            continue

        observations = concept.get("units", {}).get("USD", [])

        for item in observations:
            if item.get("form") != "10-K":
                continue

            if item.get("end") != fiscal_end:
                continue

            start = item.get("start")

            # Only full-year duration facts, not quarterly
            # or year-to-date partial-period values.
            if not start:
                continue

            try:
                days = (
                    date.fromisoformat(fiscal_end)
                    - date.fromisoformat(start)
                ).days
            except ValueError:
                continue

            if not 350 <= days <= 380:
                continue

            if item.get("val") is None:
                continue

            candidates.append({
                "value": item["val"],
                "tag": tag,
                "filed": item.get("filed", ""),
                "accession": item.get("accn", ""),
                "priority": priority,
            })

    if not candidates:
        return None

    # Prefer the original fiscal-year filing, then the
    # highest-priority accounting concept.
    original = [
        c for c in candidates
        if c["filed"].startswith(str(year))
    ]

    pool = original if original else candidates

    pool.sort(
        key=lambda x: (
            x["priority"],
            x["filed"],
        )
    )

    return pool[0]


def build_historical_data(facts):
    records = []
    source_records = []

    for year in YEARS:
        row = {"Fiscal Year": year}

        for metric, tags in TAGS.items():
            fact = get_annual_fact(facts, tags, year)

            if fact is None:
                row[metric] = None
                source_records.append({
                    "Fiscal Year": year,
                    "Metric": metric,
                    "SEC Tag": "MISSING",
                    "Filed": None,
                    "Accession": None,
                    "Value (USD Millions)": None,
                })
                continue

            value = fact["value"] / 1_000_000

            # SEC CapEx payments are generally positive
            # cash outflows. Present them as negative.
            if metric == "Capital Expenditures":
                value = -abs(value)

            row[metric] = value

            source_records.append({
                "Fiscal Year": year,
                "Metric": metric,
                "SEC Tag": fact["tag"],
                "Filed": fact["filed"],
                "Accession": fact["accession"],
                "Value (USD Millions)": value,
            })

        records.append(row)

    historical = pd.DataFrame(records)
    historical = historical.set_index("Fiscal Year")

    historical["Free Cash Flow"] = (
        historical["Operating Cash Flow"]
        + historical["Capital Expenditures"]
    )

    return historical, pd.DataFrame(source_records)


# ============================================================
# FINANCIAL RATIOS
# ============================================================

def build_ratios(df):
    ratios = pd.DataFrame(index=df.index)

    ratios["Revenue Growth"] = (
        df["Revenue"].pct_change(fill_method=None)
    )

    ratios["Operating Margin"] = (
        df["Operating Income"] / df["Revenue"]
    )

    ratios["Net Margin"] = (
        df["Net Income"] / df["Revenue"]
    )

    ratios["Free Cash Flow Margin"] = (
        df["Free Cash Flow"] / df["Revenue"]
    )

    ratios["CapEx / Revenue"] = (
        -df["Capital Expenditures"] / df["Revenue"]
    )

    ratios["D&A / Revenue"] = (
        df["Depreciation and Amortization"]
        / df["Revenue"]
    )

    return ratios


# ============================================================
# DATA QUALITY
# ============================================================

def build_quality_checks(df):
    checks = pd.DataFrame(index=df.index)

    required = [
        "Revenue",
        "Operating Income",
        "Net Income",
        "Operating Cash Flow",
        "Capital Expenditures",
    ]

    checks["Missing Core Metrics"] = (
        df[required].isna().sum(axis=1)
    )

    checks["Missing D&A"] = (
        df["Depreciation and Amortization"].isna()
    )

    checks["Revenue Positive"] = df["Revenue"] > 0

    checks["CapEx Sign Valid"] = (
        df["Capital Expenditures"] <= 0
    )

    checks["FCF Reconciliation"] = (
        df["Free Cash Flow"]
        - df["Operating Cash Flow"]
        - df["Capital Expenditures"]
    ).abs()

    checks["Status"] = checks.apply(
        lambda row: (
            "REVIEW"
            if (
                row["Missing Core Metrics"] > 0
                or row["Missing D&A"]
                or not row["Revenue Positive"]
                or not row["CapEx Sign Valid"]
                or (
                    pd.notna(row["FCF Reconciliation"])
                    and row["FCF Reconciliation"] > 0.01
                )
            )
            else "PASS"
        ),
        axis=1,
    )

    return checks


# ============================================================
# EXCEL FORMATTING
# ============================================================

def format_workbook(writer):
    for sheet in writer.book.worksheets:
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="17365D",
            )
            cell.alignment = Alignment(
                wrap_text=True,
            )

        sheet.row_dimensions[1].height = 34

        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            sheet.column_dimensions[letter].width = 25

        sheet.column_dimensions["A"].width = 32

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                 if cell.column == 1 and isinstance(cell.value, (int, float)):
                  cell.number_format = "0"
                elif isinstance(cell.value, (int, float)):
                  cell.number_format = '#,##0.00;[Red](#,##0.00)'
       

    ratios_sheet = writer.book["Financial Ratios"]

    for row in ratios_sheet.iter_rows(
        min_row=2,
        min_col=2,
    ):
        for cell in row:
            cell.number_format = "0.00%"


# ============================================================
# MAIN PROGRAM
# ============================================================

def main():
    print("Downloading Microsoft SEC company facts...")

    facts = download_company_facts()

    historical, sources = build_historical_data(facts)

    ratios = build_ratios(historical)

    checks = build_quality_checks(historical)

    metadata = pd.DataFrame([
        {
            "Field": "Company",
            "Value": COMPANY,
        },
        {
            "Field": "CIK",
            "Value": CIK,
        },
        {
            "Field": "Data Source",
            "Value": SEC_URL,
        },
        {
            "Field": "Financial Units",
            "Value": "USD Millions",
        },
        {
            "Field": "Fiscal Year End",
            "Value": "June 30",
        },
        {
            "Field": "Methodology",
            "Value": (
                "Annual US-GAAP facts from SEC Form 10-K; "
                "FCF = operating cash flow minus CapEx."
            ),
        },
    ])

    with pd.ExcelWriter(
        OUTPUT,
        engine="openpyxl",
    ) as writer:

        historical.to_excel(
            writer,
            sheet_name="Historical Financials",
        )

        ratios.to_excel(
            writer,
            sheet_name="Financial Ratios",
        )

        sources.to_excel(
            writer,
            sheet_name="SEC Source Audit",
            index=False,
        )

        checks.to_excel(
            writer,
            sheet_name="Data Quality",
        )

        metadata.to_excel(
            writer,
            sheet_name="Methodology",
            index=False,
        )

        format_workbook(writer)

    print("\nHISTORICAL FINANCIALS (USD Millions)")
    print(historical.round(2).to_string())

    print("\nDATA QUALITY")
    print(checks.to_string())

    print(f"\nExcel workbook saved to:\n{OUTPUT}")

    print(
        "\nNOTE: REVIEW rows require verification. "
        "SEC tags can differ from reported line-item "
        "definitions, especially for depreciation."
    )


if __name__ == "__main__":
    main()
