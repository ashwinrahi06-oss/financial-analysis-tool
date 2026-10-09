
"""
Microsoft Equity Research | The AI Reinvestment Test

Main project generator.

Outputs:
- Microsoft_Equity_Valuation.xlsx
- Microsoft_Forecast_Assumptions.xlsx
- Microsoft_Historical.csv
- Microsoft_Forecast.csv
- valuation_snapshot.json
- Microsoft_Research_Report.md

All fiscal-year labels are formatted as 2022, 2023, etc.
"""

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
)
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from .model_config import (
    ROOT,
    DATA,
    SEC_URL,
    FILING_URL,
    YEARS,
    FORECAST_YEARS,
    FY26_DEPRECIATION,
    FY26_DA_AND_OTHER,
    FY26_FINANCE_LEASE_ADDITIONS,
    FY26_FINANCE_LEASE_LIABILITIES,
    FY26_OPERATING_LEASE_LIABILITIES,
    RISK_FREE,
    EQUITY_RISK_PREMIUM,
    LEVERED_BETA,
    PRETAX_DEBT_COST,
    TERMINAL_GROWTH,
    SCENARIOS,
)

from .financial_data import (
    fetch_facts,
    load_financials,
)

from .forecast_assumptions import write_forecast

from .valuation_engine import (
    wacc,
    value_case,
)


# ============================================================
# CONFIGURATION
# ============================================================

OUTPUT = DATA / "Microsoft_Equity_Valuation.xlsx"

NUM_FORMAT = '#,##0.0;[Red](#,##0.0);–'
PCT_FORMAT = '0.0%;[Red](0.0%);–'
MONEY_FORMAT = '$#,##0.00;[Red]($#,##0.00);–'
YEAR_FORMAT = '0'

NAVY = "17365D"
LIGHT_BLUE = "F2F6FB"
INPUT_BLUE = "E8F2FF"
GREEN = "D9EEDC"
YELLOW = "FFF2CC"


# ============================================================
# MARKET DATA
# ============================================================

def market_data(
    price_override=None,
    shares_override=None,
):
    """
    Obtain a reference share price and share count.

    The share count is expressed in millions.

    Explicit command-line overrides take priority.
    """

    if (
        price_override is not None
        and shares_override is not None
    ):
        if price_override <= 0 or shares_override <= 0:
            raise ValueError(
                "Price and shares must be positive."
            )

        return (
            float(price_override),
            float(shares_override),
            "User-supplied market snapshot",
        )

    try:
        import yfinance as yf

        ticker = yf.Ticker("MSFT")
        info = ticker.info or {}

        price = price_override

        if price is None:
            price = (
                info.get("currentPrice")
                or info.get("regularMarketPrice")
            )

        shares = shares_override

        if shares is None:
            raw_shares = info.get("sharesOutstanding")

            if raw_shares:
                shares = float(raw_shares) / 1_000_000

        if not price or not shares:
            raise ValueError(
                "Incomplete Yahoo Finance market data."
            )

        if price <= 0 or shares <= 0:
            raise ValueError(
                "Invalid market data."
            )

        source = (
            "User-supplied / Yahoo Finance"
            if (
                price_override is not None
                or shares_override is not None
            )
            else "Yahoo Finance market snapshot"
        )

        return (
            float(price),
            float(shares),
            source,
        )

    except Exception as exc:
        raise RuntimeError(
            "Unable to retrieve a complete market snapshot. "
            "Run the script with --price and --shares-m "
            "using independently verified figures."
        ) from exc


# ============================================================
# WORKBOOK FORMATTING
# ============================================================

def decorate(
    ws,
    widths=None,
    header=True,
):
    """
    Apply professional workbook formatting.

    Fiscal-year formatting is handled separately,
    using precise cell locations rather than scanning
    all numeric values.
    """

    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B2"

    if header:
        ws.row_dimensions[1].height = 36

        for cell in ws[1]:
            cell.font = Font(
                name="Aptos",
                bold=True,
                color="FFFFFF",
                size=11,
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=NAVY,
            )

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

    for row in ws.iter_rows(min_row=2):
        for cell in row:

            if cell.data_type == "f":
                font_color = "008000"

            elif isinstance(
                cell.value,
                (int, float),
            ):
                font_color = "0000FF"

            else:
                font_color = "243247"

            cell.font = Font(
                name="Aptos",
                size=10,
                color=font_color,
            )

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

            if (
                isinstance(cell.value, (int, float))
                or cell.data_type == "f"
            ):
                cell.number_format = NUM_FORMAT

            if cell.row % 2 == 0:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=LIGHT_BLUE,
                )

    for col in range(
        1,
        ws.max_column + 1,
    ):
        letter = get_column_letter(col)

        ws.column_dimensions[letter].width = 21

    ws.column_dimensions["A"].width = 38

    for letter, width in (widths or {}).items():
        ws.column_dimensions[letter].width = width

    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def format_year_headers(
    ws,
    start_column=2,
    end_column=None,
):
    """
    Format year values in the first row.

    Example:
        2,022.0 -> 2022

    Only changes the number format.
    Does not change the underlying values.
    """

    if end_column is None:
        end_column = ws.max_column

    for column in range(
        start_column,
        end_column + 1,
    ):
        cell = ws.cell(
            row=1,
            column=column,
        )

        if isinstance(
            cell.value,
            (int, float),
        ):
            cell.number_format = YEAR_FORMAT


def highlight_input(cell):
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor=INPUT_BLUE,
    )


def highlight_output(cell):
    cell.fill = PatternFill(
        fill_type="solid",
        fgColor=GREEN,
    )


# ============================================================
# HISTORICAL FINANCIALS
# ============================================================

def build_historical_sheet(
    wb,
    history,
    da_input,
):
    ws = wb.create_sheet(
        "Historical Financials"
    )

    ws.append([
        "Metric ($M)",
        *YEARS,
    ])

    metrics = [
        "Revenue",
        "Operating Income",
        "Net Income",
        "Operating Cash Flow",
        "Cash CapEx",
        "Free Cash Flow",
    ]

    for metric in metrics:
        values = [
            float(
                history.loc[year, metric]
            )
            for year in YEARS
        ]

        ws.append([
            metric,
            *values,
        ])

    da_row = [
        "D&A model input (FY2026 only)",
    ]

    for year in YEARS:
        da_row.append(
            float(da_input)
            if year == 2026
            else None
        )

    ws.append(da_row)

    decorate(ws)

    # FIX: Display fiscal years without commas/decimals.
    format_year_headers(ws)

    ws.sheet_properties.tabColor = "2878B5"

    return ws


# ============================================================
# FORECAST INPUTS
# ============================================================

def build_forecast_sheet(
    wb,
    drivers,
):
    ws = wb.create_sheet(
        "Forecast Inputs"
    )

    ws.append([
        "Forecast Driver",
        *FORECAST_YEARS,
    ])

    driver_names = [
        "Revenue Growth",
        "EBIT Margin",
        "CapEx / Revenue",
        "D&A / Revenue",
        "Tax Rate",
        "Incremental NWC Ratio",
    ]

    driver_rows = {}

    for name in driver_names:
        values = [
            float(
                drivers.loc[i, name]
            )
            for i in range(
                len(FORECAST_YEARS)
            )
        ]

        ws.append([
            name,
            *values,
        ])

        driver_rows[name] = ws.max_row

    decorate(ws)

    # FIX: FY2027 through FY2031 display as integers.
    format_year_headers(ws)

    for row in ws.iter_rows(
        min_row=2,
        min_col=2,
    ):
        for cell in row:
            cell.number_format = PCT_FORMAT
            highlight_input(cell)

    ws.sheet_properties.tabColor = "4472C4"

    return ws, driver_rows


# ============================================================
# SCENARIO INPUTS
# ============================================================

def build_scenario_sheet(wb):
    ws = wb.create_sheet(
        "Scenario Inputs"
    )

    ws.append([
        "Scenario",
        "Revenue Growth Adjustment",
        "EBIT Margin Adjustment",
        "CapEx / Revenue Adjustment",
    ])

    scenario_rows = {}

    for name, adjustments in SCENARIOS.items():
        ws.append([
            name,
            *adjustments,
        ])

        scenario_rows[name] = ws.max_row

    decorate(
        ws,
        {
            "A": 24,
            "B": 30,
            "C": 30,
            "D": 33,
        },
    )

    for row in ws.iter_rows(
        min_row=2,
        min_col=2,
    ):
        for cell in row:
            cell.number_format = PCT_FORMAT
            highlight_input(cell)

    return ws, scenario_rows


# ============================================================
# MARKET DATA AND WACC
# ============================================================

def build_market_sheet(
    wb,
    balance,
    price,
    shares,
    quote_source,
):
    ws = wb.create_sheet(
        "Market and WACC"
    )

    ws.append([
        "Capital Structure / Cost of Capital",
        "Value",
        "Source / Methodology",
    ])

    entries = [
        (
            "Reference Price ($/share)",
            price,
            quote_source,
        ),
        (
            "Shares Outstanding (millions)",
            shares,
            quote_source,
        ),
        (
            "Cash ($M)",
            balance["Cash"],
            "SEC FY2026 Form 10-K",
        ),
        (
            "Short-Term Investments ($M)",
            (
                balance["Short-Term Investments"]
                if balance["Short-Term Investments"]
                is not None
                else 0
            ),
            "SEC FY2026 Form 10-K",
        ),
        (
            "Current Financial Debt ($M)",
            balance["Current Debt"],
            "SEC FY2026 Form 10-K",
        ),
        (
            "Long-Term Financial Debt ($M)",
            balance["Long-Term Debt"],
            "SEC FY2026 Form 10-K",
        ),
        (
            "Total Financial Debt ($M)",
            "=B6+B7",
            "Current plus long-term financial debt",
        ),
        (
            "Market Equity Capitalization ($M)",
            "=B2*B3",
            "Price multiplied by shares",
        ),
        (
            "Risk-Free Rate",
            RISK_FREE,
            "Illustrative analyst assumption",
        ),
        (
            "Equity Risk Premium",
            EQUITY_RISK_PREMIUM,
            "Illustrative analyst assumption",
        ),
        (
            "Levered Equity Beta",
            LEVERED_BETA,
            "Illustrative analyst assumption",
        ),
        (
            "Pre-Tax Cost of Debt",
            PRETAX_DEBT_COST,
            "Illustrative analyst assumption",
        ),
        (
            "Marginal Tax Rate",
            0.21,
            "Illustrative tax assumption",
        ),
        (
            "Cost of Equity",
            "=B10+B11*B12",
            "CAPM",
        ),
        (
            "WACC",
            (
                "=B9/(B9+B8)*B15"
                "+B8/(B9+B8)*B13*(1-B14)"
            ),
            "Market-value weighted cost of capital",
        ),
        (
            "Terminal Growth",
            TERMINAL_GROWTH,
            "Long-term nominal growth assumption",
        ),
        (
            "Valuation Timestamp UTC",
            datetime.now(
                timezone.utc
            ).isoformat(
                timespec="seconds"
            ),
            "Snapshot timestamp",
        ),
        (
            "Finance Lease Liabilities ($M)",
            FY26_FINANCE_LEASE_LIABILITIES,
            "FY2026 10-K lease disclosure",
        ),
        (
            "Operating Lease Liabilities ($M)",
            FY26_OPERATING_LEASE_LIABILITIES,
            "FY2026 10-K lease disclosure",
        ),
        (
            "New Finance Lease ROU Assets ($M)",
            FY26_FINANCE_LEASE_ADDITIONS,
            "Noncash lease additions",
        ),
    ]

    row_map = {}

    for name, value, source in entries:
        ws.append([
            name,
            value,
            source,
        ])

        row_map[name] = ws.max_row

    decorate(
        ws,
        {
            "A": 45,
            "B": 24,
            "C": 78,
        },
    )

    percentage_items = [
        "Risk-Free Rate",
        "Equity Risk Premium",
        "Pre-Tax Cost of Debt",
        "Marginal Tax Rate",
        "Cost of Equity",
        "WACC",
        "Terminal Growth",
    ]

    for name in percentage_items:
        row = row_map[name]

        ws[f"B{row}"].number_format = (
            PCT_FORMAT
        )

    ws["B2"].number_format = MONEY_FORMAT

    editable_items = [
        "Risk-Free Rate",
        "Equity Risk Premium",
        "Levered Equity Beta",
        "Pre-Tax Cost of Debt",
        "Marginal Tax Rate",
        "Terminal Growth",
    ]

    for name in editable_items:
        highlight_input(
            ws[f"B{row_map[name]}"]
        )

    if balance["Short-Term Investments"] is None:
        ws["C5"] = (
            "SEC XBRL value unavailable; "
            "zero is an unverified placeholder."
        )

        ws["C5"].fill = PatternFill(
            fill_type="solid",
            fgColor=YELLOW,
        )

    ws.sheet_properties.tabColor = "7030A0"

    return ws


# ============================================================
# DCF SCENARIO SHEETS
# ============================================================

def build_dcf_sheet(
    wb,
    scenario_name,
    scenario_row,
    driver_rows,
):
    ws = wb.create_sheet(
        f"{scenario_name} DCF"
    )

    ws.append([
        "Metric ($M unless stated)",
        *FORECAST_YEARS,
    ])

    labels = [
        "Revenue",
        "Revenue Growth",
        "EBIT Margin",
        "EBIT",
        "Tax Rate",
        "NOPAT",
        "D&A / Revenue",
        "D&A",
        "CapEx / Revenue",
        "Cash CapEx",
        "Incremental NWC Ratio",
        "Change NWC",
        "FCFF",
        "Discount Factor",
        "PV FCFF",
    ]

    for label in labels:
        ws.append([label])

    for i, year in enumerate(
        FORECAST_YEARS
    ):
        col = get_column_letter(i + 2)
        prev_col = get_column_letter(i + 1)

        previous_revenue = (
            "'Historical Financials'!F2"
            if i == 0
            else f"{prev_col}2"
        )

        ws[f"{col}2"] = (
            f"={previous_revenue}*(1+{col}3)"
        )

        ws[f"{col}3"] = (
            f"='Forecast Inputs'!{col}"
            f"{driver_rows['Revenue Growth']}"
            f"+'Scenario Inputs'!$B${scenario_row}"
        )

        ws[f"{col}4"] = (
            f"='Forecast Inputs'!{col}"
            f"{driver_rows['EBIT Margin']}"
            f"+'Scenario Inputs'!$C${scenario_row}"
        )

        ws[f"{col}5"] = (
            f"={col}2*{col}4"
        )

        ws[f"{col}6"] = (
            f"='Forecast Inputs'!{col}"
            f"{driver_rows['Tax Rate']}"
        )

        ws[f"{col}7"] = (
            f"={col}5*(1-{col}6)"
        )

        ws[f"{col}8"] = (
            f"='Forecast Inputs'!{col}"
            f"{driver_rows['D&A / Revenue']}"
        )

        ws[f"{col}9"] = (
            f"={col}2*{col}8"
        )

        ws[f"{col}10"] = (
            f"='Forecast Inputs'!{col}"
            f"{driver_rows['CapEx / Revenue']}"
            f"+'Scenario Inputs'!$D${scenario_row}"
        )

        ws[f"{col}11"] = (
            f"={col}2*{col}10"
        )

        ws[f"{col}12"] = (
            f"='Forecast Inputs'!{col}"
            f"{driver_rows['Incremental NWC Ratio']}"
        )

        ws[f"{col}13"] = (
            f"=({col}2-{previous_revenue})"
            f"*{col}12"
        )

        ws[f"{col}14"] = (
            f"={col}7+{col}9-{col}11-{col}13"
        )

        ws[f"{col}15"] = (
            f"=1/(1+'Market and WACC'!$B$16)"
            f"^{i+1}"
        )

        ws[f"{col}16"] = (
            f"={col}14*{col}15"
        )

    summary = {
        18: (
            "PV Forecast FCFF",
            "=SUM(B16:F16)",
        ),
        19: (
            "Terminal FCFF",
            "=F14*(1+'Market and WACC'!$B$17)",
        ),
        20: (
            "Terminal Value",
            (
                "=IF('Market and WACC'!$B$16<="
                "'Market and WACC'!$B$17,"
                "NA(),"
                "B19/('Market and WACC'!$B$16-"
                "'Market and WACC'!$B$17))"
            ),
        ),
        21: (
            "PV Terminal Value",
            "=B20*F15",
        ),
        22: (
            "Enterprise Value",
            "=B18+B21",
        ),
        23: (
            "Equity Value",
            (
                "=B22+'Market and WACC'!B4"
                "+'Market and WACC'!B5"
                "-'Market and WACC'!B8"
            ),
        ),
        24: (
            "Implied Price ($/share)",
            "=B23/'Market and WACC'!B3",
        ),
        25: (
            "Upside / Downside",
            "=B24/'Market and WACC'!B2-1",
        ),
        26: (
            "PV Terminal / EV",
            "=IFERROR(B21/B22,NA())",
        ),
    }

    for row_number, item in summary.items():
        label, formula = item

        ws[f"A{row_number}"] = label
        ws[f"B{row_number}"] = formula

    decorate(ws)

    # FIX: Forecast fiscal-year labels.
    format_year_headers(ws)

    percentage_rows = [
        3,
        4,
        6,
        8,
        10,
        12,
        15,
    ]

    for col in "BCDEF":
        for row in percentage_rows:
            ws[f"{col}{row}"].number_format = (
                PCT_FORMAT
            )

    for row in [25, 26]:
        ws[f"B{row}"].number_format = (
            PCT_FORMAT
        )

    ws["B24"].number_format = MONEY_FORMAT
    highlight_output(ws["B24"])

    if scenario_name == "Bear":
        ws.sheet_properties.tabColor = "C0504D"

    elif scenario_name == "Base":
        ws.sheet_properties.tabColor = "70AD47"

    else:
        ws.sheet_properties.tabColor = "4472C4"

    return ws


# ============================================================
# SCENARIO VALUATIONS
# ============================================================

def build_scenario_summary(wb):
    ws = wb.create_sheet(
        "Scenario Valuations",
        1,
    )

    ws.append([
        "Scenario",
        "Implied Price",
        "Reference Price",
        "Upside / Downside",
        "Enterprise Value ($M)",
        "Equity Value ($M)",
        "PV Terminal Weight",
    ])

    for name in SCENARIOS:
        sheet = f"{name} DCF"

        ws.append([
            name,
            f"='{sheet}'!B24",
            "='Market and WACC'!B2",
            f"='{sheet}'!B25",
            f"='{sheet}'!B22",
            f"='{sheet}'!B23",
            f"='{sheet}'!B26",
        ])

    decorate(
        ws,
        {
            "A": 20,
            "B": 23,
            "C": 23,
            "D": 25,
            "E": 29,
            "F": 26,
            "G": 26,
        },
    )

    for row in range(2, 5):
        for col in "BC":
            ws[f"{col}{row}"].number_format = (
                MONEY_FORMAT
            )

        for col in "DG":
            ws[f"{col}{row}"].number_format = (
                PCT_FORMAT
            )

    return ws


# ============================================================
# SENSITIVITY ANALYSIS
# ============================================================

def build_sensitivity_sheet(wb):
    ws = wb.create_sheet("Sensitivity")

    growth_offsets = [
        -0.01,
        -0.005,
        0,
        0.005,
        0.01,
    ]

    wacc_offsets = [
        -0.02,
        -0.01,
        0,
        0.01,
        0.02,
    ]

    ws["A1"] = "WACC / Terminal Growth"

    for col_number, offset in enumerate(
        growth_offsets,
        start=2,
    ):
        col = get_column_letter(col_number)

        ws[f"{col}1"] = (
            f"='Market and WACC'!$B$17"
            f"{offset:+.3f}"
        )

    for row_number, offset in enumerate(
        wacc_offsets,
        start=2,
    ):
        ws[f"A{row_number}"] = (
            f"='Market and WACC'!$B$16"
            f"{offset:+.3f}"
        )

        for col_number in range(2, 7):
            col = get_column_letter(col_number)

            pv_terms = []

            for i in range(5):
                forecast_col = get_column_letter(
                    i + 2
                )

                pv_terms.append(
                    f"'Base DCF'!{forecast_col}14"
                    f"/(1+$A{row_number})^{i+1}"
                )

            pv_forecast = "+".join(pv_terms)

            pv_terminal = (
                f"('Base DCF'!F14*(1+{col}$1)"
                f"/($A{row_number}-{col}$1))"
                f"/(1+$A{row_number})^5"
            )

            equity_bridge = (
                "'Market and WACC'!$B$4"
                "+'Market and WACC'!$B$5"
                "-'Market and WACC'!$B$8"
            )

            ws[
                f"{col}{row_number}"
            ] = (
                f"=IF($A{row_number}<={col}$1,"
                f"NA(),"
                f"(({pv_forecast})+{pv_terminal}"
                f"+{equity_bridge})"
                f"/'Market and WACC'!$B$3)"
            )

    decorate(
        ws,
        {"A": 34},
    )

    for col in "BCDEF":
        ws[f"{col}1"].number_format = (
            PCT_FORMAT
        )

        for row in range(2, 7):
            ws[f"{col}{row}"].number_format = (
                MONEY_FORMAT
            )

    for row in range(2, 7):
        ws[f"A{row}"].number_format = (
            PCT_FORMAT
        )

    highlight_output(ws["D4"])

    return ws


# ============================================================
# SEC SOURCE AUDIT
# ============================================================

def build_audit_sheet(
    wb,
    audit,
):
    ws = wb.create_sheet(
        "SEC Source Audit"
    )

    ws.append(
        list(audit.columns)
    )

    for record in audit.itertuples(
        index=False,
        name=None,
    ):
        ws.append(list(record))

    decorate(
        ws,
        {
            "A": 19,
            "B": 33,
            "C": 59,
            "D": 20,
            "E": 33,
            "F": 24,
        },
    )

    # FIX: The first column contains fiscal years.
    for row in range(
        2,
        ws.max_row + 1,
    ):
        ws[f"A{row}"].number_format = (
            YEAR_FORMAT
        )

    return ws


# ============================================================
# ACCOUNTING RECONCILIATION
# ============================================================

def build_accounting_sheet(
    wb,
    history,
    da_input,
):
    ws = wb.create_sheet(
        "Accounting Reconciliation"
    )

    ws.append([
        "Accounting Item",
        "USD Millions",
        "Interpretation",
    ])

    records = [
        (
            "FY2026 Cash PP&E Purchases",
            float(
                history.loc[2026, "Cash CapEx"]
            ),
            "Cash investing outflow used in FCFF",
        ),
        (
            "FY2026 Reported Depreciation",
            FY26_DEPRECIATION,
            "Not necessarily total D&A",
        ),
        (
            "FY2026 D&A and Other",
            FY26_DA_AND_OTHER,
            "Broader adjustment, not pure D&A",
        ),
        (
            "FY2026 D&A Model Input",
            da_input,
            "Depreciation-only starting assumption",
        ),
        (
            "New Finance Lease ROU Assets",
            FY26_FINANCE_LEASE_ADDITIONS,
            "Noncash infrastructure investment",
        ),
        (
            "Finance Lease Liabilities",
            FY26_FINANCE_LEASE_LIABILITIES,
            "Separately disclosed lease obligation",
        ),
        (
            "Operating Lease Liabilities",
            FY26_OPERATING_LEASE_LIABILITIES,
            "Separately disclosed lease obligation",
        ),
    ]

    for record in records:
        ws.append(list(record))

    decorate(
        ws,
        {
            "A": 46,
            "B": 23,
            "C": 83,
        },
    )

    for row in ws.iter_rows(min_row=2):
        if (
            row[0].value
            == "FY2026 D&A Model Input"
        ):
            for cell in row:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=YELLOW,
                )

    return ws


# ============================================================
# METHODOLOGY
# ============================================================

def build_methodology_sheet(wb):
    ws = wb.create_sheet("Methodology")

    ws.append([
        "Model Methodology and Limitations"
    ])

    notes = [
        "Historical financials are sourced from SEC Company Facts.",
        "Forecasts are analyst assumptions, not Microsoft guidance.",
        "FCFF = EBIT x (1 - tax) + D&A - Cash CapEx - Change in NWC.",
        "WACC is calculated using CAPM and market-value capital structure.",
        "D&A uses a depreciation-only starting assumption.",
        "The D&A starting value is not verified total depreciation and amortization.",
        "Cash CapEx excludes noncash finance lease additions.",
        "A fully lease-adjusted model requires consistent cash flow and debt treatment.",
        "The equity bridge includes cash and short-term investments less financial debt.",
        "Missing short-term investments are provisionally treated as zero and flagged.",
        "Terminal value may represent a substantial portion of enterprise value.",
        "Changing WACC or terminal growth will recalculate the Excel DCF.",
        "Excel formulas recalculate when opened in a compatible spreadsheet application.",
        "SEC Source: " + SEC_URL,
        "Annual Filing: " + FILING_URL,
    ]

    for note in notes:
        ws.append([note])

    decorate(
        ws,
        {"A": 125},
    )

    return ws


# ============================================================
# INVESTMENT DASHBOARD
# ============================================================

def build_dashboard(wb):
    ws = wb.create_sheet(
        "Investment Dashboard",
        0,
    )

    rows = [
        (
            "MICROSOFT | THE AI REINVESTMENT TEST",
            None,
        ),
        (
            "Key Valuation Output",
            "Formula-Linked Result",
        ),
        (
            "Base Implied Price",
            "='Base DCF'!B24",
        ),
        (
            "Reference Share Price",
            "='Market and WACC'!B2",
        ),
        (
            "Base Upside / Downside",
            "='Base DCF'!B25",
        ),
        (
            "Calculated WACC",
            "='Market and WACC'!B16",
        ),
        (
            "Terminal Growth",
            "='Market and WACC'!B17",
        ),
        (
            "Terminal Value Weight",
            "='Base DCF'!B26",
        ),
        (
            "D&A Accounting Status",
            "Depreciation-only input; total D&A not reconciled",
        ),
        (
            "Research Question",
            "Can AI infrastructure investment generate sufficient returns?",
        ),
        (
            "Annual Report",
            FILING_URL,
        ),
    ]

    for row in rows:
        ws.append(list(row))

    decorate(
        ws,
        {
            "A": 55,
            "B": 95,
        },
    )

    ws["A1"].font = Font(
        name="Aptos Display",
        size=16,
        bold=True,
        color=NAVY,
    )

    for row in [3, 4]:
        ws[f"B{row}"].number_format = (
            MONEY_FORMAT
        )

    for row in [5, 6, 7, 8]:
        ws[f"B{row}"].number_format = (
            PCT_FORMAT
        )

    ws["B9"].fill = PatternFill(
        fill_type="solid",
        fgColor=YELLOW,
    )

    ws.sheet_properties.tabColor = "17365D"

    return ws


# ============================================================
# COMPLETE EXCEL WORKBOOK
# ============================================================

def create_workbook(
    history,
    balance,
    audit,
    drivers,
    price,
    shares,
    quote_source,
    da_input,
):
    wb = Workbook()

    wb.remove(wb.active)

    wb.calculation = CalcProperties(
        calcMode="auto",
        fullCalcOnLoad=True,
        forceFullCalc=True,
    )

    build_historical_sheet(
        wb,
        history,
        da_input,
    )

    _, driver_rows = build_forecast_sheet(
        wb,
        drivers,
    )

    _, scenario_rows = build_scenario_sheet(
        wb
    )

    build_market_sheet(
        wb,
        balance,
        price,
        shares,
        quote_source,
    )

    for scenario_name in SCENARIOS:
        build_dcf_sheet(
            wb,
            scenario_name,
            scenario_rows[scenario_name],
            driver_rows,
        )

    build_scenario_summary(wb)

    build_sensitivity_sheet(wb)

    build_audit_sheet(
        wb,
        audit,
    )

    build_accounting_sheet(
        wb,
        history,
        da_input,
    )

    build_methodology_sheet(wb)

    build_dashboard(wb)

    return wb


# ============================================================
# EQUITY RESEARCH REPORT
# ============================================================

def report_markdown(
    history,
    drivers,
    results,
    price,
    shares,
    quote_source,
    wacc_value,
    da_input,
):
    timestamp = datetime.now(
        timezone.utc
    ).strftime("%Y-%m-%d %H:%M UTC")

    historical_rows = []

    for year in YEARS:
        historical_rows.append(
            f"| {year} | "
            f"{history.loc[year, 'Revenue']:,.0f} | "
            f"{history.loc[year, 'Operating Income']:,.0f} | "
            f"{history.loc[year, 'Cash CapEx']:,.0f} | "
            f"{history.loc[year, 'Free Cash Flow']:,.0f} |"
        )

    forecast_rows = []

    for _, row in drivers.iterrows():
        forecast_rows.append(
            f"| {int(row['Fiscal Year'])} | "
            f"{row['Revenue Growth']:.1%} | "
            f"{row['EBIT Margin']:.1%} | "
            f"{row['CapEx / Revenue']:.1%} | "
            f"{row['D&A / Revenue']:.1%} |"
        )

    scenario_rows = []

    for name, result in results.items():
        scenario_rows.append(
            f"| {name} | "
            f"${result['Implied Price']:,.2f} | "
            f"{result['Upside']:+.1%} | "
            f"{result['PV Terminal Weight']:.1%} |"
        )

    historical_table = "\n".join(
        historical_rows
    )

    forecast_table = "\n".join(
        forecast_rows
    )

    scenario_table = "\n".join(
        scenario_rows
    )

    base = results["Base"]

    return f"""# Microsoft (MSFT): The AI Reinvestment Test

**Research Date:** {timestamp}

**Reference Price:** ${price:,.2f}

**Shares Outstanding:** {shares:,.1f} million

**Market Data Source:** {quote_source}

## Executive Summary

This research examines whether Microsoft's investments
in cloud computing and artificial intelligence can
generate sufficient future cash flows to justify
the capital required.

The central question is not simply whether Microsoft's
revenue will grow. It is whether incremental operating
profits will translate into sustainable free cash flow
after infrastructure investment.

## Investment Thesis

Microsoft benefits from an established enterprise
software ecosystem, Azure cloud infrastructure,
and multiple channels for AI monetization.

However, substantial capital spending creates
uncertainty about near-term free cash flow and
long-term returns on invested capital.

This model evaluates that tradeoff using three
explicit operating scenarios.

## Historical Financial Performance

USD millions.

| Fiscal Year | Revenue | Operating Income | Cash CapEx | Free Cash Flow |
|---|---:|---:|---:|---:|
{historical_table}

## Forecast Assumptions

| Fiscal Year | Revenue Growth | EBIT Margin | CapEx / Revenue | D&A / Revenue |
|---|---:|---:|---:|---:|
{forecast_table}

These forecasts are analyst assumptions,
not Microsoft management guidance.

## Discount Rate

**Calculated WACC:** {wacc_value:.2%}

**Terminal Growth:** {TERMINAL_GROWTH:.2%}

WACC is based on illustrative CAPM inputs and
the modeled market-value capital structure.

## Valuation Results

| Scenario | Implied Price | Upside / Downside | PV Terminal Weight |
|---|---:|---:|---:|
{scenario_table}

**Base Implied Price:** ${base['Implied Price']:,.2f}

**Reference Price:** ${price:,.2f}

**Base Upside / Downside:** {base['Upside']:+.1%}

The valuation is conditional on the assumptions.
It is not calibrated to match the stock price.

## AI Infrastructure Reinvestment

FY2026 cash CapEx totaled approximately
${history.loc[2026, 'Cash CapEx']:,.0f} million.

The valuation explicitly forecasts cash capital
expenditures as a percentage of revenue.

It also discloses finance-lease additions separately
because cash CapEx does not capture all infrastructure
investment commitments.

## Accounting Reconciliation

- Reported FY2026 depreciation:
  ${FY26_DEPRECIATION:,.0f} million.
- Model D&A starting input:
  ${da_input:,.0f} million.
- Depreciation, amortization and other:
  ${FY26_DA_AND_OTHER:,.0f} million.
- New finance-lease ROU assets:
  ${FY26_FINANCE_LEASE_ADDITIONS:,.0f} million.
- Finance-lease liabilities:
  ${FY26_FINANCE_LEASE_LIABILITIES:,.0f} million.

The D&A model input is a depreciation-only
starting assumption and is not a verified
standalone total D&A figure.

The current valuation does not fully capitalize
lease-related investments or lease financing.

## Bull-Case Catalysts

- Faster Azure and AI revenue monetization.
- Improving utilization of computing infrastructure.
- Higher operating margins.
- Moderating cash capital expenditure intensity.
- Lower required returns.

## Bear-Case Risks

- Slower AI monetization.
- Continued high infrastructure spending.
- Shorter-than-expected computing asset lives.
- Lower operating leverage.
- Higher discount rates.

## Research Limitations

The model uses simplified assumptions for
working capital, taxation, depreciation,
capital expenditure trajectories, and
cost of capital.

Short-term investments require independent
verification if the SEC XBRL field is missing.

The terminal value can represent a large
portion of estimated enterprise value.

The model should not be represented as a
fully reconciled institutional research valuation
until these assumptions are validated.

## Sources

- [Microsoft FY2026 Form 10-K]({FILING_URL})
- [SEC Company Facts]({SEC_URL})
- [Microsoft Investor Relations](https://www.microsoft.com/investor)

*Educational equity research case study.*
"""


# ============================================================
# MAIN PROJECT GENERATOR
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description=(
            "Build Microsoft equity research "
            "and DCF valuation files."
        )
    )

    parser.add_argument(
        "--price",
        type=float,
        help="Reference share price",
    )

    parser.add_argument(
        "--shares-m",
        type=float,
        help="Shares outstanding in millions",
    )

    parser.add_argument(
        "--da",
        type=float,
        default=FY26_DEPRECIATION,
        help=(
            "FY2026 D&A model input "
            "in USD millions"
        ),
    )

    parser.add_argument(
        "--facts-json",
        type=Path,
        help=(
            "Optional offline SEC "
            "Company Facts JSON"
        ),
    )

    args = parser.parse_args()

    if args.da <= 0:
        raise ValueError(
            "D&A must be positive."
        )

    # --------------------------------------------------------
    # LOAD HISTORICAL FINANCIALS
    # --------------------------------------------------------

    if args.facts_json:
        facts = json.loads(
            args.facts_json.read_text(
                encoding="utf-8"
            )
        )

    else:
        facts = fetch_facts()

    history, balance, audit = load_financials(
        facts
    )

    # --------------------------------------------------------
    # LOAD MARKET SNAPSHOT
    # --------------------------------------------------------

    price, shares, quote_source = market_data(
        args.price,
        args.shares_m,
    )

    # --------------------------------------------------------
    # GENERATE FORECAST ASSUMPTIONS
    # --------------------------------------------------------

    da_status = (
        "DEPRECIATION-ONLY FLOOR"
        if args.da == FY26_DEPRECIATION
        else "USER INPUT / VERIFY"
    )

    forecast, forecast_path = write_forecast(
        history,
        args.da,
        da_status,
    )

    # --------------------------------------------------------
    # CALCULATE WACC
    # --------------------------------------------------------

    debt = (
        balance["Current Debt"]
        + balance["Long-Term Debt"]
    )

    investments = (
        balance["Short-Term Investments"]
        if balance["Short-Term Investments"]
        is not None
        else 0
    )

    rate = wacc(
        RISK_FREE,
        EQUITY_RISK_PREMIUM,
        LEVERED_BETA,
        PRETAX_DEBT_COST,
        0.21,
        price * shares,
        debt,
    )

    if rate <= TERMINAL_GROWTH:
        raise ValueError(
            "WACC must exceed terminal growth."
        )

    # --------------------------------------------------------
    # CALCULATE SCENARIO VALUATIONS
    # --------------------------------------------------------

    records = forecast.to_dict(
        orient="records"
    )

    results = {}

    for name, adjustments in SCENARIOS.items():
        results[name] = value_case(
            float(
                history.loc[2026, "Revenue"]
            ),
            records,
            adjustments,
            rate,
            TERMINAL_GROWTH,
            balance["Cash"],
            investments,
            debt,
            shares,
            price,
        )

    # --------------------------------------------------------
    # GENERATE EXCEL MODEL
    # --------------------------------------------------------

    workbook = create_workbook(
        history,
        balance,
        audit,
        forecast,
        price,
        shares,
        quote_source,
        args.da,
    )

    workbook.save(OUTPUT)

    # --------------------------------------------------------
    # GENERATE CSV FILES
    # --------------------------------------------------------

    historical_csv = (
        DATA / "Microsoft_Historical.csv"
    )

    forecast_csv = (
        DATA / "Microsoft_Forecast.csv"
    )

    history.to_csv(
        historical_csv
    )

    forecast.to_csv(
        forecast_csv,
        index=False,
    )

    # --------------------------------------------------------
    # GENERATE VALUATION SNAPSHOT
    # --------------------------------------------------------

    snapshot = {
        "as_of_utc": datetime.now(
            timezone.utc
        ).isoformat(),
        "quote_source": quote_source,
        "price": price,
        "shares_m": shares,
        "cash_m": balance["Cash"],
        "short_term_investments_m": investments,
        "debt_m": debt,
        "wacc": rate,
        "terminal_growth": TERMINAL_GROWTH,
        "da_input_m": args.da,
        "scenarios": results,
    }

    snapshot_path = (
        DATA / "valuation_snapshot.json"
    )

    snapshot_path.write_text(
        json.dumps(
            snapshot,
            indent=2,
        ),
        encoding="utf-8",
    )

    # --------------------------------------------------------
    # GENERATE RESEARCH REPORT
    # --------------------------------------------------------

    report_path = (
        ROOT / "Microsoft_Research_Report.md"
    )

    report = report_markdown(
        history,
        forecast,
        results,
        price,
        shares,
        quote_source,
        rate,
        args.da,
    )

    report_path.write_text(
        report,
        encoding="utf-8",
    )

    # --------------------------------------------------------
    # RESULTS
    # --------------------------------------------------------

    print()
    print("=" * 65)
    print("MICROSOFT EQUITY RESEARCH")
    print("THE AI REINVESTMENT TEST")
    print("=" * 65)

    print(
        f"Reference Price: ${price:,.2f}"
    )

    print(
        f"WACC: {rate:.2%}"
    )

    print(
        f"Terminal Growth: {TERMINAL_GROWTH:.2%}"
    )

    print()

    for name, result in results.items():
        print(
            f"{name:5s} | "
            f"Implied Price: "
            f"${result['Implied Price']:,.2f} | "
            f"Upside: {result['Upside']:+.1%}"
        )

    print()
    print("FILES GENERATED:")
    print(OUTPUT)
    print(forecast_path)
    print(historical_csv)
    print(forecast_csv)
    print(snapshot_path)
    print(report_path)

    print()
    print(
        "Fiscal-year formatting applied "
        "to historical, forecast, DCF, "
        "and SEC audit worksheets."
    )

    print(
        "WARNING: D&A and selected market/accounting "
        "assumptions remain provisional."
    )

    print("=" * 65)


if __name__ == "__main__":
    main()
