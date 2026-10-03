import yfinance as yf
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.chart import LineChart, Reference

def get_value(statement, row_name, year):
    if row_name in statement.index:
        return statement.loc[row_name, year]
    else:
        return None

def format_money(value):
    if value is None:
        return "N/A"
    elif abs(value) >= 1_000_000_000:
        return f"${value / 1_000_000_000:.2f}B"
    elif abs(value) >= 1_000_000:
        return f"${value / 1_000_000:.2f}M"
    else:
        return f"${value:,.2f}"

ticker = input("Enter a stock ticker: ").upper()

stock = yf.Ticker(ticker)

income_statement = stock.financials
balance_sheet = stock.balance_sheet

print()
print("AVAILABLE FINANCIAL YEARS")
print("-------------------------")
print(income_statement.columns)

historical_data = []

for year in income_statement.columns:
    historical_revenue = income_statement.loc["Total Revenue", year]
    historical_net_income = income_statement.loc["Net Income", year]

    historical_data.append({
        "Year": year.year,
        "Revenue": historical_revenue,
        "Net Income": historical_net_income
    })

    historical_df = pd.DataFrame(historical_data)

    historical_df["Revenue Growth (%)"] = (
    historical_df["Revenue"].pct_change(periods=-1) * 100
).round(2)

    historical_df["Net Margin (%)"] = (
    historical_df["Net Income"] / historical_df["Revenue"] * 100
).round(2)

print()
print("HISTORICAL PERFORMANCE")
print("-------------------------")
print(historical_df.to_string(index=False))

print(
        year.year,
        format_money(historical_revenue),
        format_money(historical_net_income)
    )

if income_statement.empty or balance_sheet.empty:
    print("Financial data could not be found for this ticker.")
    exit()

latest_year = income_statement.columns[0]
previous_year = income_statement.columns[1]
latest_balance_year = balance_sheet.columns[0]

revenue = income_statement.loc["Total Revenue", latest_year]
previous_revenue = income_statement.loc["Total Revenue", previous_year]
net_income = income_statement.loc["Net Income", latest_year]

total_assets = balance_sheet.loc["Total Assets", latest_balance_year]
total_equity = balance_sheet.loc["Stockholders Equity", latest_balance_year]
current_assets = balance_sheet.loc["Current Assets", latest_balance_year]
current_liabilities = balance_sheet.loc["Current Liabilities", latest_balance_year]
total_debt = get_value(
    balance_sheet,
    "Total Debt",
    latest_balance_year
)
net_margin = round((net_income / revenue) * 100, 2)

roa = round((net_income / total_assets) * 100, 2)

roe = round((net_income / total_equity) * 100, 2)

current_ratio = round(current_assets / current_liabilities, 2)

if total_debt is not None:
    debt_to_equity = round(total_debt / total_equity, 2)
else:
    debt_to_equity = "N/A"

revenue_growth = round(
    ((revenue - previous_revenue) / previous_revenue) * 100,
    2
)

report = {
    "Metric": [
        "Ticker",
        "Year",
        "Revenue",
        "Previous Revenue",
        "Revenue Growth",
        "Net Income",
        "Net Profit Margin",
        "Total Assets",
        "Stockholders Equity",
        "Return on Assets",
        "Return on Equity",
        "Current Assets",
        "Current Liabilities",
        "Current Ratio",
        "Total Debt",
        "Debt-to-Equity"
    ],

    "Value": [
        ticker,
        latest_year.year,
        format_money(revenue),
        format_money(previous_revenue),
        str(revenue_growth) + "%",
        format_money(net_income),
        str(net_margin) + "%",
        format_money(total_assets),
        format_money(total_equity),
        str(roa) + "%",
        str(roe) + "%",
        format_money(current_assets),
        format_money(current_liabilities),
        current_ratio,
        format_money(total_debt),
        debt_to_equity
    ]
}

report_df = pd.DataFrame(report)

print()
print(ticker, "FINANCIAL DATA")
print("-------------------------")
print("Year:", latest_year.year)
print("Revenue:", format_money(revenue))
print("Net Income:", format_money(net_income))
print("Net Profit Margin:", net_margin, "%")
print("Total Assets:", format_money(total_assets))
print("Stockholders Equity:", format_money(total_equity))
print("Return on Assets:", roa, "%")
print("Return on Equity:", roe, "%")
print("Current Assets:", format_money(current_assets))
print("Current Liabilities:", format_money(current_liabilities))
print("Current Ratio:", current_ratio)
print("Total Debt:", format_money(total_debt))
print("Debt-to-Equity:", debt_to_equity)
print("Previous Revenue:", format_money(previous_revenue))
print("Revenue Growth:", revenue_growth, "%")

print()
print("FINANCIAL REPORT")
print("-------------------------")
print(report_df.to_string(index=False))

file_name = ticker + "_financial_report.xlsx"

with pd.ExcelWriter(file_name, engine="openpyxl") as writer:
    report_df.to_excel(
        writer,
        sheet_name="Financial Summary",
        index=False
    )

    historical_df.to_excel(
        writer,
        sheet_name="Historical Performance",
        index=False
    )
workbook = load_workbook(file_name)

worksheet = workbook["Financial Summary"]
historical_worksheet = workbook["Historical Performance"]

for cell in worksheet[1]:
    cell.font = Font(bold=True)

for column in worksheet.columns:
    max_length = 0
    column_letter = column[0].column_letter

    for cell in column:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

    worksheet.column_dimensions[column_letter].width = max_length + 2

worksheet.freeze_panes = "A2"
worksheet.auto_filter.ref = worksheet.dimensions

for cell in historical_worksheet[1]:
    cell.font = Font(bold=True)

for column in historical_worksheet.columns:
    max_length = 0
    column_letter = column[0].column_letter

    for cell in column:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

    historical_worksheet.column_dimensions[column_letter].width = max_length + 2

historical_worksheet.column_dimensions["B"].width = 22
historical_worksheet.column_dimensions["C"].width = 22

for row in historical_worksheet.iter_rows(min_row=2):
    # Revenue
    row[1].number_format = '$#,##0'

    # Net Income
    row[2].number_format = '$#,##0'

    # Revenue Growth
    if row[3].value is not None:
        row[3].value = row[3].value / 100
        row[3].number_format = '0.00%'

    # Net Margin
    if row[4].value is not None:
        row[4].value = row[4].value / 100
        row[4].number_format = '0.00%'

historical_worksheet.freeze_panes = "A2"

historical_worksheet.auto_filter.ref = historical_worksheet.dimensions

# Create historical financial performance chart
chart = LineChart()

chart.title = "Revenue and Net Income"
chart.y_axis.title = "Amount ($)"
chart.x_axis.title = "Year"

# Revenue and Net Income data
data = Reference(
    historical_worksheet,
    min_col=2,
    max_col=3,
    min_row=1,
    max_row=historical_worksheet.max_row
)

# Years for the horizontal axis
years = Reference(
    historical_worksheet,
    min_col=1,
    min_row=2,
    max_row=historical_worksheet.max_row
)

chart.add_data(data, titles_from_data=True)
chart.set_categories(years)

chart.height = 8
chart.width = 14

historical_worksheet.add_chart(chart, "G2")

workbook.save(file_name)

print()
print("Excel report saved as:", file_name)
