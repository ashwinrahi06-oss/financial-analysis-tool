
# ============================================================
# FINANCIAL INTELLIGENCE DASHBOARD
# Finance + MIS Portfolio Project
# ============================================================

import io
import math
import openpyxl
import pandas as pd
import streamlit as st
import yfinance as yf

from case_studies.microsoft_valuation.equity_research_page import render as equity_research_page
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from dcf_page import show_dcf_page

from company_comparison import (
    get_company_overview,
    build_comparison,
    format_value,
    METRIC_ORDER,
    MONEY_METRICS,
    PERCENT_METRICS,
    MULTIPLE_METRICS,
)


from pathlib import Path
import sys

MICROSOFT_DIR = (
    Path(__file__).resolve().parent
    / "case_studies"
    / "microsoft_valuation"
)





# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Financial Intelligence Dashboard",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="expanded",
)

NAVY = "#17365D"

st.markdown(
    """
    <style>
    .block-container {
        padding-top: 2rem;
        max-width: 1400px;
    }

    div[data-testid="stMetric"] {
        background-color: rgba(128, 128, 128, 0.07);
        padding: 16px;
        border-radius: 10px;
        border: 1px solid rgba(128, 128, 128, 0.15);
    }

    div[data-testid="stMetricLabel"] {
        font-weight: 600;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# DATA RETRIEVAL
# ============================================================

@st.cache_data(ttl=3600, show_spinner=False)
def load_company(ticker):
    return get_company_overview(ticker)


@st.cache_data(ttl=3600, show_spinner=False)
def load_historical_financials(ticker):

    stock = yf.Ticker(ticker)

    income = stock.financials
    cashflow = stock.cashflow

    if income.empty:
        raise ValueError(
            f"No annual financial statements available for {ticker}."
        )

    rows = []

    for year in sorted(income.columns):

        def statement_value(statement, metric):
            if statement.empty:
                return None

            if metric not in statement.index:
                return None

            if year not in statement.columns:
                return None

            value = statement.loc[metric, year]

            if isinstance(value, pd.Series):
                value = value.iloc[0]

            if pd.isna(value):
                return None

            return float(value)

        revenue = statement_value(
            income, "Total Revenue"
        )

        operating_income = statement_value(
            income, "Operating Income"
        )

        net_income = statement_value(
            income, "Net Income"
        )

        operating_cash_flow = statement_value(
            cashflow, "Operating Cash Flow"
        )

        capex = statement_value(
            cashflow, "Capital Expenditure"
        )

        fcf = None

        if (
            operating_cash_flow is not None
            and capex is not None
        ):
            fcf = operating_cash_flow + capex

        operating_margin = (
            operating_income / revenue * 100
            if revenue not in (None, 0)
            and operating_income is not None
            else None
        )

        net_margin = (
            net_income / revenue * 100
            if revenue not in (None, 0)
            and net_income is not None
            else None
        )

        rows.append({
            "Fiscal Year": year.strftime("%Y-%m-%d"),
            "Revenue": revenue,
            "Operating Income": operating_income,
            "Net Income": net_income,
            "Operating Cash Flow": operating_cash_flow,
            "Capital Expenditures": capex,
            "Free Cash Flow": fcf,
            "Operating Margin (%)": operating_margin,
            "Net Margin (%)": net_margin,
        })

    df = pd.DataFrame(rows)

    if df.empty:
        return df

    df["Revenue Growth (%)"] = (
        df["Revenue"].pct_change(fill_method=None) * 100
    )

    df["FCF Margin (%)"] = (
        df["Free Cash Flow"] / df["Revenue"] * 100
    )

    return df


# ============================================================
# FORMAT HELPERS
# ============================================================

def safe_format(metric, value):
    if value is None:
        return "N/A"

    try:
        if pd.isna(value):
            return "N/A"
    except (TypeError, ValueError):
        pass

    return format_value(metric, value)


def to_billions(df, columns):
    result = df[["Fiscal Year"] + columns].copy()

    for column in columns:
        result[column] = result[column] / 1e9

    return result.set_index("Fiscal Year")


def excel_number_format(metric):
    if metric in MONEY_METRICS:
        return '"$"#,##0;[Red]("$"#,##0)'

    if metric in PERCENT_METRICS:
        return '0.00"%"'

    if metric in MULTIPLE_METRICS:
        return '0.00"x"'

    return "General"


def style_workbook(workbook):

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "B2"

        worksheet.row_dimensions[1].height = 28

        for cell in worksheet[1]:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="17365D",
            )
            cell.font = Font(
                color="FFFFFF",
                bold=True,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for column in worksheet.columns:
            letter = column[0].column_letter
            worksheet.column_dimensions[letter].width = 23

        worksheet.column_dimensions["A"].width = 31

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if cell.row % 2 == 0:
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F2F6FA",
                    )


# ============================================================
# EXCEL EXPORT
# ============================================================

def build_excel_report(
    company,
    historical_df=None,
    other_company=None,
    comparison_df=None,
    quality_df=None,
):

    output = io.BytesIO()

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.title = "Financial Summary"

    ticker = company["ticker"]

    worksheet.append([
        "Financial Metric",
        ticker,
    ])

    for metric in METRIC_ORDER:
        worksheet.append([
            metric,
            company["metrics"].get(metric),
        ])

    for row in range(2, worksheet.max_row + 1):
        metric = worksheet.cell(row, 1).value
        worksheet.cell(row, 2).number_format = (
            excel_number_format(metric)
        )

    if historical_df is not None and not historical_df.empty:

        historical_sheet = workbook.create_sheet(
            "Historical Performance"
        )

        for row in [historical_df.columns.tolist()] + (
            historical_df.where(
                pd.notna(historical_df), None
            ).values.tolist()
        ):
            historical_sheet.append(list(row))

        for column_number, column_name in enumerate(
            historical_df.columns,
            start=1,
        ):
            if column_name == "Fiscal Year":
                continue

            for row_number in range(
                2, historical_sheet.max_row + 1
            ):
                historical_sheet.cell(
                    row_number, column_number
                ).number_format = excel_number_format(
                    column_name
                )

    if comparison_df is not None:

        comparison_sheet = workbook.create_sheet(
            "Company Comparison"
        )

        comparison_sheet.append([
            "Financial Metric",
            *comparison_df.columns.tolist(),
        ])

        for metric in comparison_df.index:

            values = [
                comparison_df.loc[metric, ticker_name]
                for ticker_name in comparison_df.columns
            ]

            comparison_sheet.append([
                metric,
                *values,
            ])

        for row_number in range(
            2, comparison_sheet.max_row + 1
        ):
            metric = comparison_sheet.cell(
                row_number, 1
            ).value

            for col_number in range(
                2, comparison_sheet.max_column + 1
            ):
                comparison_sheet.cell(
                    row_number, col_number
                ).number_format = excel_number_format(
                    metric
                )

    if quality_df is not None:

        quality_sheet = workbook.create_sheet(
            "Data Quality"
        )

        quality_sheet.append(
            quality_df.columns.tolist()
        )

        for row in quality_df.itertuples(
            index=False,
            name=None,
        ):
            quality_sheet.append(list(row))

        quality_sheet.column_dimensions["D"].width = 65

    methodology = workbook.create_sheet("Methodology")

    methodology.append(["Topic", "Explanation"])

    notes = [
        (
            "Source",
            "Yahoo Finance data retrieved through yfinance",
        ),
        (
            "Annual Statements",
            "Financial metrics are based on annual reporting periods",
        ),
        (
            "Revenue Growth",
            "Current revenue divided by prior revenue minus one",
        ),
        (
            "Free Cash Flow",
            "Operating cash flow plus signed capital expenditures",
        ),
        (
            "Percentages",
            "Stored as percentage-point values, e.g. 25 means 25%",
        ),
        (
            "Fiscal Year Comparison",
            "Companies may have different fiscal year-end dates",
        ),
        (
            "Valuation",
            "Market-based valuation metrics may use a different as-of date",
        ),
        (
            "Missing Values",
            "Unavailable metrics are left blank or shown as N/A",
        ),
        (
            "Limitations",
            "Automated data is subject to source errors and accounting differences",
        ),
    ]

    for note in notes:
        methodology.append(note)

    methodology.column_dimensions["B"].width = 95

    style_workbook(workbook)

    workbook.save(output)
    output.seek(0)

    return output.getvalue()


# ============================================================
# COMPANY OVERVIEW
# ============================================================

def show_company_overview(company):

    metrics = company["metrics"]

    st.subheader(
        f"{metrics['Company']} ({company['ticker']})"
    )

    st.caption(
        f"{metrics['Sector']} · {metrics['Industry']} "
        f"· Fiscal year ending {metrics['Fiscal Year']}"
    )

    columns = st.columns(4)

    overview_metrics = [
        ("Market Cap", "Market Cap"),
        ("P/E Ratio", "P/E Ratio"),
        ("EV/EBITDA", "EV/EBITDA"),
        ("Revenue", "Annual Revenue"),
    ]

    for column, (metric, label) in zip(
        columns, overview_metrics
    ):
        with column:
            st.metric(
                label,
                safe_format(
                    metric,
                    metrics.get(metric),
                ),
            )


def show_financial_metrics(company):

    metrics = company["metrics"]

    st.subheader("Financial Performance")

    first_row = st.columns(3)

    items = [
        ("Revenue Growth (%)", "Revenue Growth"),
        ("Operating Margin (%)", "Operating Margin"),
        ("Net Margin (%)", "Net Margin"),
    ]

    for column, (metric, label) in zip(
        first_row, items
    ):
        with column:
            st.metric(
                label,
                safe_format(metric, metrics.get(metric)),
            )

    second_row = st.columns(3)

    items = [
        ("Operating Cash Flow", "Operating Cash Flow"),
        ("Free Cash Flow", "Free Cash Flow"),
        ("FCF Margin (%)", "FCF Margin"),
    ]

    for column, (metric, label) in zip(
        second_row, items
    ):
        with column:
            st.metric(
                label,
                safe_format(metric, metrics.get(metric)),
            )


# ============================================================
# HISTORICAL CHARTS
# ============================================================

def show_historical_charts(historical_df):

    if historical_df.empty:
        st.info("Historical financial data is unavailable.")
        return

    st.subheader("Historical Financial Performance")

    st.caption(
        "Annual financial statement data. Monetary charts use USD billions."
    )

    money_columns = [
        "Revenue",
        "Operating Income",
        "Net Income",
    ]

    st.markdown("**Revenue and Profitability**")

    st.line_chart(
        to_billions(historical_df, money_columns),
        y_label="USD (Billions)",
        x_label="Fiscal Year",
    )

    cash_columns = [
        "Operating Cash Flow",
        "Free Cash Flow",
    ]

    st.markdown("**Cash Flow Performance**")

    st.line_chart(
        to_billions(historical_df, cash_columns),
        y_label="USD (Billions)",
        x_label="Fiscal Year",
    )

    margin_columns = [
        "Operating Margin (%)",
        "Net Margin (%)",
        "FCF Margin (%)",
    ]

    st.markdown("**Profitability Margins**")

    st.line_chart(
        historical_df[
            ["Fiscal Year"] + margin_columns
        ].set_index("Fiscal Year"),
        y_label="Percent (%)",
        x_label="Fiscal Year",
    )

    with st.expander("View Historical Financial Data"):

        display_df = historical_df.copy()

        st.dataframe(
            display_df,
            use_container_width=True,
            hide_index=True,
        )


# ============================================================
# FINANCIAL HEALTH INDICATORS
# ============================================================

def show_financial_health(company):

    metrics = company["metrics"]

    st.subheader("Financial Health Indicators")

    revenue_growth = metrics.get("Revenue Growth (%)")
    net_margin = metrics.get("Net Margin (%)")
    fcf = metrics.get("Free Cash Flow")

    indicators = []

    if revenue_growth is None:
        growth_status = "Unavailable"
    elif revenue_growth > 0:
        growth_status = "Positive"
    elif revenue_growth < 0:
        growth_status = "Negative"
    else:
        growth_status = "Flat"

    indicators.append({
        "Category": "Growth",
        "Status": growth_status,
        "Metric": "Revenue Growth",
        "Value": safe_format(
            "Revenue Growth (%)",
            revenue_growth,
        ),
    })

    if net_margin is None:
        margin_status = "Unavailable"
    elif net_margin > 0:
        margin_status = "Positive"
    elif net_margin < 0:
        margin_status = "Negative"
    else:
        margin_status = "Break-even"

    indicators.append({
        "Category": "Profitability",
        "Status": margin_status,
        "Metric": "Net Margin",
        "Value": safe_format(
            "Net Margin (%)",
            net_margin,
        ),
    })

    if fcf is None:
        cash_status = "Unavailable"
    elif fcf > 0:
        cash_status = "Positive"
    elif fcf < 0:
        cash_status = "Negative"
    else:
        cash_status = "Neutral"

    indicators.append({
        "Category": "Cash Flow",
        "Status": cash_status,
        "Metric": "Free Cash Flow",
        "Value": safe_format(
            "Free Cash Flow",
            fcf,
        ),
    })

    st.dataframe(
        pd.DataFrame(indicators),
        use_container_width=True,
        hide_index=True,
    )

    st.caption(
        "These indicators describe individual metrics, not an "
        "overall investment recommendation."
    )


# ============================================================
# SINGLE-COMPANY ANALYSIS PAGE
# ============================================================

def single_company_page():

    st.header("Company Financial Analysis")

    ticker = st.text_input(
        "Enter a stock ticker",
        value="AAPL",
        key="single_ticker",
    ).upper().strip()

    if st.button(
        "Analyze Company",
        type="primary",
        key="single_analyze",
    ):

        if not ticker:
            st.warning("Please enter a ticker.")
            return

        with st.spinner("Analyzing financial statements..."):

            try:
                company = load_company(ticker)

                historical_df = load_historical_financials(
                    ticker
                )

                st.session_state["single_result"] = (
                    ticker,
                    company,
                    historical_df,
                )

            except Exception as error:
                st.error(f"Analysis failed: {error}")
                return

    result = st.session_state.get("single_result")

    if result is None:
        st.info("Enter a ticker and select Analyze Company.")
        return

    result_ticker, company, historical_df = result

    if result_ticker != ticker:
        st.info(
            "The displayed results are from your previous analysis. "
            "Select Analyze Company to update them."
        )

    st.divider()

    show_company_overview(company)

    st.divider()

    show_financial_metrics(company)

    st.divider()

    show_historical_charts(historical_df)

    st.divider()

    show_financial_health(company)

    st.divider()

    excel_bytes = build_excel_report(
        company,
        historical_df=historical_df,
    )

    st.download_button(
        "Download Financial Analysis Excel Report",
        data=excel_bytes,
        file_name=f"{result_ticker}_financial_dashboard.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        type="primary",
    )


# ============================================================
# COMPANY COMPARISON PAGE
# ============================================================

def company_comparison_page():

    st.header("Company Comparison")

    col1, col2 = st.columns(2)

    with col1:
        ticker_a = st.text_input(
            "First Company",
            value="AAPL",
            key="comparison_a",
        ).upper().strip()

    with col2:
        ticker_b = st.text_input(
            "Second Company",
            value="MSFT",
            key="comparison_b",
        ).upper().strip()

    if st.button(
        "Compare Companies",
        type="primary",
        key="compare_button",
    ):

        if not ticker_a or not ticker_b:
            st.warning("Please enter both tickers.")
            return

        if ticker_a == ticker_b:
            st.warning("Please select different companies.")
            return

        with st.spinner("Comparing companies..."):

            try:

                company_a = load_company(ticker_a)
                company_b = load_company(ticker_b)

                comparison_df, formatted_df, quality_df = (
                    build_comparison(
                        company_a,
                        company_b,
                    )
                )

                st.session_state["comparison_result"] = (
                    company_a,
                    company_b,
                    comparison_df,
                    formatted_df,
                    quality_df,
                )

            except Exception as error:
                st.error(f"Comparison failed: {error}")
                return

    result = st.session_state.get("comparison_result")

    if result is None:
        st.info(
            "Enter two tickers and select Compare Companies."
        )
        return

    (
        company_a,
        company_b,
        comparison_df,
        formatted_df,
        quality_df,
    ) = result

    first = company_a["ticker"]
    second = company_b["ticker"]

    st.divider()

    st.subheader(f"{first} vs {second}")

    st.dataframe(
        formatted_df,
        use_container_width=True,
    )

    if (
        company_a["latest_period"]
        != company_b["latest_period"]
    ):
        st.warning(
            "These companies have different fiscal year-end dates: "
            f"{first} ({company_a['metrics']['Fiscal Year']}) "
            f"and {second} ({company_b['metrics']['Fiscal Year']}). "
            "The comparison does not align identical calendar periods."
        )

    if (
        company_a["metrics"]["Sector"]
        != company_b["metrics"]["Sector"]
    ):
        st.info(
            "These companies operate in different sectors. "
            "Industry-specific accounting and business models "
            "can make direct metric comparisons misleading."
        )

    st.divider()

    st.subheader("Financial Performance Comparison")

    financial_metrics = [
        "Revenue",
        "Operating Income",
        "Net Income",
        "Operating Cash Flow",
        "Free Cash Flow",
    ]

    financial_chart = comparison_df.loc[
        financial_metrics
    ].copy()

    financial_chart = financial_chart.apply(
        pd.to_numeric,
        errors="coerce",
    ) / 1e9

    st.bar_chart(
        financial_chart,
        x_label="Financial Metric",
        y_label="USD (Billions)",
    )

    st.subheader("Growth and Profitability Comparison")

    percent_metrics = [
        "Revenue Growth (%)",
        "Operating Margin (%)",
        "Net Margin (%)",
        "FCF Margin (%)",
    ]

    percent_chart = comparison_df.loc[
        percent_metrics
    ].copy()

    percent_chart = percent_chart.apply(
        pd.to_numeric,
        errors="coerce",
    )

    st.bar_chart(
        percent_chart,
        x_label="Financial Metric",
        y_label="Percent (%)",
    )

    st.divider()

    with st.expander("View Data Quality and Limitations"):
        st.dataframe(
            quality_df,
            use_container_width=True,
            hide_index=True,
        )

    excel_bytes = build_excel_report(
        company_a,
        other_company=company_b,
        comparison_df=comparison_df,
        quality_df=quality_df,
    )

    st.download_button(
        "Download Company Comparison Excel Report",
        data=excel_bytes,
        file_name=f"{first}_vs_{second}_dashboard.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        type="primary",
    )


# ============================================================
# SIDEBAR NAVIGATION
# ============================================================

with st.sidebar:

    st.title("Financial Intelligence")

    st.caption("Finance + MIS Analytics Platform")

    st.divider()

    page = st.radio(
    "Navigation",
    [
        "Company Analysis",
        "Company Comparison",
        "Microsoft Equity Research",
        "About the Project",
    ],
)

    st.divider()

    st.caption(
        "Financial data provided by Yahoo Finance "
        "through yfinance."
    )


# ============================================================
# MAIN DASHBOARD
# ============================================================

st.title("Financial Intelligence Dashboard")

st.write(
    "An interactive platform for evaluating public "
    "companies through financial statement analysis, "
    "historical performance, and peer comparisons."
)

if page == "Company Analysis":

    single_company_page()

elif page == "Company Comparison":

    company_comparison_page()

elif page == "Microsoft Equity Research":

    equity_research_page()

elif page == "About the Project":

    st.header("About the Project")

    st.write(
        """
        The Financial Intelligence Dashboard is an
        interactive financial analysis and equity
        research platform built using Python.

        It combines automated financial statement
        analysis, company comparisons, discounted
        cash flow valuation, and investment research.

        ### Core Capabilities

        **1. Company Financial Analysis**
        - Automated financial statement retrieval
        - Revenue growth and profitability analysis
        - Operating cash flow and free cash flow analysis
        - Historical financial performance charts
        - Downloadable Excel financial reports

        **2. Company Comparison**
        - Side-by-side public company comparisons
        - Financial performance benchmarking
        - Profitability and growth comparisons
        - Data quality and reporting-period transparency

        **3. Microsoft Equity Research**
        - Historical SEC financial statement analysis
        - Five-year financial forecasts
        - Discounted cash flow valuation
        - Weighted average cost of capital analysis
        - Bear, base, and bull scenarios
        - Valuation sensitivity analysis
        - AI infrastructure reinvestment analysis
        - Downloadable valuation workbook

        ### Technologies

        Python, Pandas, Streamlit, yfinance,
        openpyxl, Plotly, and SEC Company Facts.

        ### Project Objective

        Demonstrate the application of financial
        modeling, corporate valuation, data analytics,
        and financial technology to real-world
        investment research.

        ### Methodology

        Historical data is retrieved from financial
        data providers and SEC filings.

        Forecast assumptions are analyst estimates,
        not company guidance.

        The valuation models are intended for
        educational research and are not investment
        recommendations.
        """
    )

