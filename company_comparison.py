
# ============================================================
# COMPANY COMPARISON ENGINE
# Financial Analysis Tool | Finance Portfolio Project
# ============================================================

import math
import pandas as pd
import yfinance as yf

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

MONEY_METRICS = {
    "Market Cap",
    "Revenue",
    "Operating Income",
    "Net Income",
    "Operating Cash Flow",
    "Capital Expenditures",
    "Free Cash Flow"
}

PERCENT_METRICS = {
    "Revenue Growth (%)",
    "Operating Margin (%)",
    "Net Margin (%)",
    "FCF Margin (%)"
}

MULTIPLE_METRICS = {
    "P/E Ratio",
    "EV/EBITDA"
}

METRIC_ORDER = [
    "Company",
    "Sector",
    "Industry",
    "Fiscal Year",
    "Market Cap",
    "P/E Ratio",
    "EV/EBITDA",
    "Revenue",
    "Revenue Growth (%)",
    "Operating Income",
    "Operating Margin (%)",
    "Net Income",
    "Net Margin (%)",
    "Operating Cash Flow",
    "Capital Expenditures",
    "Free Cash Flow",
    "FCF Margin (%)"
]

BANKING_INDUSTRIES = {
    "banks - diversified",
    "banks - regional"
}


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean_number(value):
    """Convert financial values to floats safely."""

    if value is None:
        return None

    try:
        number = float(value)

        if not math.isfinite(number):
            return None

        return number

    except (TypeError, ValueError):
        return None


def get_statement_value(statement, metric, period):
    """Retrieve a value from a financial statement."""

    if statement.empty or period is None:
        return None

    if metric not in statement.index:
        return None

    if period not in statement.columns:
        return None

    value = statement.loc[metric, period]

    if isinstance(value, pd.Series):
        value = value.iloc[0]

    return clean_number(value)


def calculate_ratio(numerator, denominator):
    """Calculate a percentage without dividing by zero."""

    if numerator is None or denominator in (None, 0):
        return None

    return numerator / denominator * 100


def format_value(metric, value):
    """Format values for readable terminal output."""

    if value is None:
        return "N/A"

    if isinstance(value, str):
        return value

    if metric in MONEY_METRICS:

        if abs(value) >= 1_000_000_000_000:
            return f"${value / 1_000_000_000_000:,.2f}T"

        if abs(value) >= 1_000_000_000:
            return f"${value / 1_000_000_000:,.2f}B"

        if abs(value) >= 1_000_000:
            return f"${value / 1_000_000:,.2f}M"

        return f"${value:,.2f}"

    if metric in PERCENT_METRICS:
        return f"{value:.2f}%"

    if metric in MULTIPLE_METRICS:
        return f"{value:.2f}x"

    return str(value)


def get_metric_note(metric, company):
    """Explain data availability and methodology limitations."""

    if metric in company["excluded_metrics"]:
        return "Not assessed using conventional banking methodology."

    if company["metrics"].get(metric) is None:
        return "Unavailable from source data or insufficient inputs."

    return "Available"


# ============================================================
# RETRIEVE AND ANALYZE COMPANY
# ============================================================

def get_company_overview(ticker):

    ticker = ticker.upper().strip()

    if not ticker:
        raise ValueError("Ticker cannot be empty.")

    stock = yf.Ticker(ticker)

    info = stock.info or {}

    income_statement = stock.financials
    cash_flow_statement = stock.cashflow

    if income_statement.empty:
        raise ValueError(
            f"No annual income statement found for {ticker}."
        )

    # Sort fiscal periods from newest to oldest
    income_periods = sorted(
        income_statement.columns,
        reverse=True
    )

    latest_year = income_periods[0]

    previous_year = (
        income_periods[1]
        if len(income_periods) > 1
        else None
    )

    revenue = get_statement_value(
        income_statement,
        "Total Revenue",
        latest_year
    )

    previous_revenue = get_statement_value(
        income_statement,
        "Total Revenue",
        previous_year
    )

    operating_income = get_statement_value(
        income_statement,
        "Operating Income",
        latest_year
    )

    net_income = get_statement_value(
        income_statement,
        "Net Income",
        latest_year
    )

    revenue_growth = None

    if (
        revenue is not None
        and previous_revenue is not None
        and previous_revenue != 0
    ):
        revenue_growth = (
            revenue / previous_revenue - 1
        ) * 100

    operating_margin = calculate_ratio(
        operating_income,
        revenue
    )

    net_margin = calculate_ratio(
        net_income,
        revenue
    )

    # Cash flow must match the income statement period
    operating_cash_flow = get_statement_value(
        cash_flow_statement,
        "Operating Cash Flow",
        latest_year
    )

    capital_expenditures = get_statement_value(
        cash_flow_statement,
        "Capital Expenditure",
        latest_year
    )

    free_cash_flow = None

    if (
        operating_cash_flow is not None
        and capital_expenditures is not None
    ):
        # yfinance generally represents CapEx as negative
        free_cash_flow = (
            operating_cash_flow + capital_expenditures
        )

    fcf_margin = calculate_ratio(
        free_cash_flow,
        revenue
    )

    sector = info.get("sector") or "N/A"
    industry = info.get("industry") or "N/A"

    is_bank = (
        sector == "Financial Services"
        and industry.lower() in BANKING_INDUSTRIES
    )

    excluded_metrics = set()

    if is_bank:
        excluded_metrics.update({
            "Operating Margin (%)",
            "Free Cash Flow",
            "FCF Margin (%)"
        })

    metrics = {
        "Company": info.get("longName") or ticker,
        "Sector": sector,
        "Industry": industry,
        "Fiscal Year": latest_year.strftime("%Y-%m-%d"),
        "Market Cap": clean_number(info.get("marketCap")),
        "P/E Ratio": clean_number(info.get("trailingPE")),
        "EV/EBITDA": clean_number(
            info.get("enterpriseToEbitda")
        ),
        "Revenue": revenue,
        "Revenue Growth (%)": revenue_growth,
        "Operating Income": operating_income,
        "Operating Margin (%)": operating_margin,
        "Net Income": net_income,
        "Net Margin (%)": net_margin,
        "Operating Cash Flow": operating_cash_flow,
        "Capital Expenditures": capital_expenditures,
        "Free Cash Flow": free_cash_flow,
        "FCF Margin (%)": fcf_margin
    }

    # Do not present excluded metrics as standard
    # operating-company comparison measures.
    for metric in excluded_metrics:
        metrics[metric] = None

    return {
        "ticker": ticker,
        "metrics": metrics,
        "excluded_metrics": excluded_metrics,
        "latest_period": latest_year,
        "previous_period": previous_year,
        "is_bank": is_bank
    }


# ============================================================
# BUILD COMPARISON TABLES
# ============================================================

def build_comparison(company_a, company_b):

    ticker_a = company_a["ticker"]
    ticker_b = company_b["ticker"]

    if ticker_a == ticker_b:
        raise ValueError(
            "Choose two different stock tickers."
        )

    comparison_df = pd.DataFrame({
        ticker_a: company_a["metrics"],
        ticker_b: company_b["metrics"]
    })

    comparison_df = comparison_df.reindex(METRIC_ORDER)
    comparison_df.index.name = "Financial Metric"

    formatted_df = comparison_df.copy()

    for metric in formatted_df.index:
        for ticker in formatted_df.columns:

            formatted_df.loc[metric, ticker] = format_value(
                metric,
                comparison_df.loc[metric, ticker]
            )

    quality_rows = []

    for company in [company_a, company_b]:

        for metric in METRIC_ORDER:

            if metric in {
                "Company",
                "Sector",
                "Industry",
                "Fiscal Year"
            }:
                continue

            quality_rows.append({
                "Ticker": company["ticker"],
                "Metric": metric,
                "Status": (
                    "Not Assessed"
                    if metric in company["excluded_metrics"]
                    else "Missing"
                    if company["metrics"].get(metric) is None
                    else "Available"
                ),
                "Explanation": get_metric_note(
                    metric,
                    company
                )
            })

    quality_df = pd.DataFrame(quality_rows)

    return comparison_df, formatted_df, quality_df


# ============================================================
# EXCEL STYLING
# ============================================================

NAVY = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F6FA"
WHITE = "FFFFFF"


def style_header(worksheet):

    for cell in worksheet[1]:

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=NAVY
        )

        cell.font = Font(
            bold=True,
            color=WHITE,
            size=11
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    worksheet.row_dimensions[1].height = 28


def style_comparison_sheet(worksheet):

    style_header(worksheet)

    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 28

    worksheet.freeze_panes = "B2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for row in range(2, worksheet.max_row + 1):

        metric = worksheet.cell(row, 1).value

        worksheet.cell(row, 1).font = Font(
            bold=True
        )

        if row % 2 == 0:
            for cell in worksheet[row]:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=LIGHT_GRAY
                )

        for col in [2, 3]:

            cell = worksheet.cell(row, col)

            if metric in MONEY_METRICS:
                cell.number_format = (
                    '"$"#,##0;[Red]("$"#,##0)'
                )

            elif metric in PERCENT_METRICS:
                # Values are stored as percentage points
                cell.number_format = '0.00"%"'

            elif metric in MULTIPLE_METRICS:
                cell.number_format = '0.00"x"'

            if cell.value is None:
                cell.value = "N/A"


def style_quality_sheet(worksheet):

    style_header(worksheet)

    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 20
    worksheet.column_dimensions["D"].width = 66

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    colors = {
        "Available": "E2F0D9",
        "Missing": "FCE4D6",
        "Not Assessed": "E7E6E6"
    }

    for row in worksheet.iter_rows(min_row=2):

        status_cell = row[2]

        status_cell.fill = PatternFill(
            fill_type="solid",
            fgColor=colors.get(
                status_cell.value,
                WHITE
            )
        )

        row[3].alignment = Alignment(
            wrap_text=True,
            vertical="center"
        )

        worksheet.row_dimensions[row[0].row].height = 28


# ============================================================
# CREATE COMPARISON CHARTS
# ============================================================

def create_comparison_charts(
    workbook,
    comparison_df,
    ticker_a,
    ticker_b
):

    chart_sheet = workbook.create_sheet(
        "Comparison Charts"
    )

    chart_sheet["A1"] = "Financial Performance Comparison"

    chart_sheet["A1"].font = Font(
        bold=True,
        size=16,
        color=NAVY
    )

    chart_sheet["A2"] = (
        "Annual financial metrics; companies may "
        "have different fiscal year-end dates."
    )

    chart_sheet.column_dimensions["A"].width = 28
    chart_sheet.column_dimensions["B"].width = 20
    chart_sheet.column_dimensions["C"].width = 20

    # Chart source data: values converted to billions
    money_chart_metrics = [
        "Revenue",
        "Operating Income",
        "Net Income",
        "Operating Cash Flow",
        "Free Cash Flow"
    ]

    chart_sheet.append([])
    chart_sheet.append([
        "Metric",
        ticker_a,
        ticker_b
    ])

    for metric in money_chart_metrics:

        values = [
            comparison_df.loc[metric, ticker]
            for ticker in [ticker_a, ticker_b]
        ]

        chart_sheet.append([
            metric,
            *[
                value / 1_000_000_000
                if isinstance(value, (int, float))
                and math.isfinite(value)
                else None
                for value in values
            ]
        ])

    for row in chart_sheet.iter_rows(
        min_row=4,
        max_row=9,
        min_col=2,
        max_col=3
    ):
        for cell in row:
            cell.number_format = "0.0"

    money_chart = BarChart()

    money_chart.type = "col"
    money_chart.grouping = "clustered"
    money_chart.overlap = 0

    money_chart.title = "Financial Performance"
    money_chart.y_axis.title = "USD (Billions)"
    money_chart.x_axis.title = "Financial Metric"

    money_chart.width = 23
    money_chart.height = 12
    money_chart.style = 10

    money_data = Reference(
        chart_sheet,
        min_col=2,
        max_col=3,
        min_row=4,
        max_row=9
    )

    money_categories = Reference(
        chart_sheet,
        min_col=1,
        min_row=5,
        max_row=9
    )

    money_chart.add_data(
        money_data,
        titles_from_data=True
    )

    money_chart.set_categories(
        money_categories
    )

    money_chart.legend.position = "b"

    chart_sheet.add_chart(
        money_chart,
        "E4"
    )

    # Percentage comparison table
    chart_sheet["A12"] = "Metric"
    chart_sheet["B12"] = ticker_a
    chart_sheet["C12"] = ticker_b

    percent_chart_metrics = [
        "Revenue Growth (%)",
        "Operating Margin (%)",
        "Net Margin (%)",
        "FCF Margin (%)"
    ]

    for row_number, metric in enumerate(
        percent_chart_metrics,
        start=13
    ):

        chart_sheet.cell(row_number, 1, metric)

        for col_number, ticker in [
            (2, ticker_a),
            (3, ticker_b)
        ]:

            value = comparison_df.loc[metric, ticker]

            cell = chart_sheet.cell(
                row_number,
                col_number
            )

            if isinstance(value, (int, float)):
                if math.isfinite(value):
                    cell.value = value

            cell.number_format = '0.0"%"'

    percent_chart = BarChart()

    percent_chart.type = "col"
    percent_chart.grouping = "clustered"
    percent_chart.overlap = 0

    percent_chart.title = "Growth and Profitability"
    percent_chart.y_axis.title = "Percent (%)"
    percent_chart.x_axis.title = "Financial Metric"

    percent_chart.width = 23
    percent_chart.height = 12
    percent_chart.style = 10

    percent_data = Reference(
        chart_sheet,
        min_col=2,
        max_col=3,
        min_row=12,
        max_row=16
    )

    percent_categories = Reference(
        chart_sheet,
        min_col=1,
        min_row=13,
        max_row=16
    )

    percent_chart.add_data(
        percent_data,
        titles_from_data=True
    )

    percent_chart.set_categories(
        percent_categories
    )

    percent_chart.legend.position = "b"

    chart_sheet.add_chart(
        percent_chart,
        "E28"
    )

    for row_number in [4, 12]:
        for cell in chart_sheet[row_number][:3]:

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=NAVY
            )

            cell.font = Font(
                bold=True,
                color=WHITE
            )


# ============================================================
# EXPORT EXCEL REPORT
# ============================================================

def export_comparison(
    company_a,
    company_b,
    comparison_df,
    quality_df
):

    ticker_a = company_a["ticker"]
    ticker_b = company_b["ticker"]

    file_name = (
        f"{ticker_a}_vs_{ticker_b}_comparison.xlsx"
    )

    # Export raw numeric data, not formatted text
    with pd.ExcelWriter(
        file_name,
        engine="openpyxl"
    ) as writer:

        comparison_df.to_excel(
            writer,
            sheet_name="Company Comparison",
            index=True
        )

        quality_df.to_excel(
            writer,
            sheet_name="Data Quality",
            index=False
        )

    workbook = load_workbook(file_name)

    comparison_sheet = workbook[
        "Company Comparison"
    ]

    quality_sheet = workbook[
        "Data Quality"
    ]

    style_comparison_sheet(
        comparison_sheet
    )

    style_quality_sheet(
        quality_sheet
    )

    create_comparison_charts(
        workbook,
        comparison_df,
        ticker_a,
        ticker_b
    )

    # Methodology worksheet
    methodology = workbook.create_sheet(
        "Methodology"
    )

    notes = [
        ("Topic", "Explanation"),
        (
            "Data Source",
            "Yahoo Finance via yfinance"
        ),
        (
            "Reporting Frequency",
            "Annual financial statements"
        ),
        (
            "Revenue Growth",
            "Current annual revenue / prior annual revenue - 1"
        ),
        (
            "Operating Margin",
            "Operating income / revenue"
        ),
        (
            "Net Margin",
            "Net income / revenue"
        ),
        (
            "Free Cash Flow",
            "Operating cash flow + signed capital expenditures"
        ),
        (
            "FCF Margin",
            "Free cash flow / revenue"
        ),
        (
            "Valuation",
            "Market cap and valuation multiples reflect source availability and may use a different as-of date from annual financial statements."
        ),
        (
            "Fiscal Periods",
            "Each company uses its latest available annual income statement period. Periods may differ between companies."
        ),
        (
            "Industry Limitations",
            "Certain conventional operating-company metrics are not assessed for banks."
        ),
        (
            "Missing Data",
            "N/A indicates unavailable or insufficient source information."
        ),
        (
            "Purpose",
            "Educational financial analysis; not investment advice."
        )
    ]

    for row in notes:
        methodology.append(row)

    methodology.column_dimensions["A"].width = 27
    methodology.column_dimensions["B"].width = 105

    style_header(methodology)

    for row in methodology.iter_rows(min_row=2):
        row[1].alignment = Alignment(
            wrap_text=True,
            vertical="center"
        )

        methodology.row_dimensions[
            row[0].row
        ].height = 32

    workbook.save(file_name)

    return file_name


# ============================================================
# MAIN PROGRAM
# ============================================================

def main():

    print()
    print("=" * 60)
    print("COMPANY COMPARISON ENGINE")
    print("=" * 60)

    ticker_a = input(
        "Enter first stock ticker: "
    ).upper().strip()

    ticker_b = input(
        "Enter second stock ticker: "
    ).upper().strip()

    if not ticker_a or not ticker_b:
        print("Error: Both tickers are required.")
        return

    if ticker_a == ticker_b:
        print("Error: Enter two different companies.")
        return

    print()
    print("Retrieving financial data...")

    try:

        company_a = get_company_overview(ticker_a)
        company_b = get_company_overview(ticker_b)

        comparison_df, formatted_df, quality_df = (
            build_comparison(
                company_a,
                company_b
            )
        )

        print()
        print("=" * 60)
        print("COMPANY FINANCIAL COMPARISON")
        print("=" * 60)

        print(
            formatted_df.to_string()
        )

        # Reporting-period warning
        if (
            company_a["latest_period"]
            != company_b["latest_period"]
        ):
            print()
            print("REPORTING PERIOD NOTICE:")
            print(
                f"{ticker_a}: "
                f"{company_a['metrics']['Fiscal Year']}"
            )
            print(
                f"{ticker_b}: "
                f"{company_b['metrics']['Fiscal Year']}"
            )
            print(
                "The companies have different fiscal "
                "year-end dates. Interpret comparisons "
                "with this limitation in mind."
            )

        # Industry warning
        if (
            company_a["metrics"]["Sector"]
            != company_b["metrics"]["Sector"]
        ):
            print()
            print("INDUSTRY COMPARISON NOTICE:")
            print(
                "These companies operate in different "
                "sectors. Differences in margins, "
                "capital intensity, and valuation "
                "may reflect their business models."
            )

        file_name = export_comparison(
            company_a,
            company_b,
            comparison_df,
            quality_df
        )

        print()
        print("=" * 60)
        print("REPORT GENERATED SUCCESSFULLY")
        print("=" * 60)
        print(f"File: {file_name}")
        print("Worksheets:")
        print(" - Company Comparison")
        print(" - Data Quality")
        print(" - Comparison Charts")
        print(" - Methodology")

    except Exception as error:

        print()
        print("An error occurred:")
        print(error)


if __name__ == "__main__":
    main()
